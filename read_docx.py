from docx import Document
import sys

doc = Document(sys.argv[1])
for p in doc.paragraphs:
    print(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                print(f"[Table Cell] {p.text}")
